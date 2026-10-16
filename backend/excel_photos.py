import logging
import re
from io import BytesIO
from pathlib import Path
from urllib.request import Request, urlopen

from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

try:
    from PIL import Image as PILImage, ImageOps
    PILLOW_AVAILABLE = True
except ImportError:
    PILLOW_AVAILABLE = False


PHOTO_COLUMN = 8
PHOTO_COLUMN_WIDTH = 45
PHOTO_ROW_HEIGHT = 185

DISPLAY_BOX = (180, 240)
EMBED_MAX = (400, 533)
JPEG_QUALITY = 80

VISIT_ID_PATTERN = re.compile(
    r"visit[_\-\s]?(\d+)",
    re.IGNORECASE
)

GALENIUM_BASE_URL = "https://outlet.galenium.com/"


def build_photo_index(image_dir):
    index = {
        "name": {},
        "stem": {},
        "id": {}
    }

    image_dir = Path(image_dir)

    if not image_dir.exists():
        return index

    for path in image_dir.rglob("*"):
        if not path.is_file():
            continue

        name = path.name.lower()
        stem = path.stem.lower()

        index["name"][name] = path
        index["stem"][stem] = path

        match = VISIT_ID_PATTERN.search(path.name)
        if match:
            visit_id = match.group(1)
            index["id"][visit_id] = path

    return index


def find_photo(value, index, visit_id=None):
    if not value:
        return None

    value = str(value).strip()

    # Bersihkan query string
    value = value.split("?")[0]

    # Normalisasi slash
    value = value.replace("\\", "/")

    # Ambil nama file terakhir
    filename = value.rstrip("/").split("/")[-1]

    if not filename:
        return None

    filename_lower = filename.lower()
    stem_lower = Path(filename).stem.lower()

    # 1. Berdasarkan nama file
    if filename_lower in index["name"]:
        return index["name"][filename_lower]

    # 2. Berdasarkan nama tanpa extension
    if stem_lower in index["stem"]:
        return index["stem"][stem_lower]

    # 3. Berdasarkan visit ID
    if visit_id is not None:
        visit_id_str = str(visit_id)

        if visit_id_str in index["id"]:
            return index["id"][visit_id_str]

    return None


def build_external_url(value):
    if not value:
        return None

    value = str(value).strip()

    if not value:
        return None

    # Kalau sudah URL lengkap
    if value.startswith("http://") or value.startswith("https://"):
        return value

    # Kalau masih berupa:
    # images/abc.jpeg
    # /images/abc.jpeg
    value = value.lstrip("/")

    return GALENIUM_BASE_URL + value


def _prepare_photo(path):
    if not PILLOW_AVAILABLE:
        raise RuntimeError("Pillow belum terinstall.")

    with PILImage.open(path) as image:
        image = ImageOps.exif_transpose(image)

        if image.mode != "RGB":
            image = image.convert("RGB")

        # Jika foto landscape, putar menjadi portrait
        if image.width > image.height:
            image = image.rotate(90, expand=True)

        # Pastikan ukuran maksimal portrait
        image.thumbnail(EMBED_MAX, PILImage.LANCZOS)

        output = BytesIO()

        image.save(
            output,
            format="JPEG",
            quality=JPEG_QUALITY,
            optimize=True
        )

        output.seek(0)

        return output


def _prepare_photo_bytes(image_bytes):
    if not PILLOW_AVAILABLE:
        raise RuntimeError("Pillow belum terinstall.")

    with PILImage.open(BytesIO(image_bytes)) as image:
        image = ImageOps.exif_transpose(image)

        if image.mode != "RGB":
            image = image.convert("RGB")

        # Jika foto landscape, putar menjadi portrait
        if image.width > image.height:
            image = image.rotate(90, expand=True)

        # Pastikan ukuran maksimal portrait
        image.thumbnail(EMBED_MAX, PILImage.LANCZOS)

        output = BytesIO()

        image.save(
            output,
            format="JPEG",
            quality=JPEG_QUALITY,
            optimize=True
        )

        output.seek(0)

        return output


def _prepare_photo_bytes(image_bytes):
    if not PILLOW_AVAILABLE:
        raise RuntimeError("Pillow belum terinstall.")

    with PILImage.open(BytesIO(image_bytes)) as image:
        image = ImageOps.exif_transpose(image)

        if image.mode != "RGB":
            image = image.convert("RGB")

        # Landscape -> portrait
        if image.width > image.height:
            image = image.rotate(90, expand=True)

        image.thumbnail(EMBED_MAX, PILImage.LANCZOS)

        output = BytesIO()

        image.save(
            output,
            format="JPEG",
            quality=JPEG_QUALITY,
            optimize=True
        )

        output.seek(0)

        return output


def _display_size(width, height):
    if width <= 0 or height <= 0:
        return DISPLAY_BOX

    scale = min(
        DISPLAY_BOX[0] / width,
        DISPLAY_BOX[1] / height
    )

    return (
        int(width * scale),
        int(height * scale)
    )


def download_external_photo(url):
    request = Request(
        url,
        headers={
            "User-Agent": "Mozilla/5.0"
        }
    )

    with urlopen(request, timeout=15) as response:
        content_type = response.headers.get("Content-Type", "")

        image_bytes = response.read()

    # Pastikan yang didownload memang gambar
    if content_type and not content_type.lower().startswith("image/"):
        raise ValueError(
            f"Response bukan gambar: {content_type}"
        )

    if not image_bytes:
        raise ValueError("File gambar kosong.")

    return image_bytes


def embed_photos(
    worksheet,
    rows,
    image_dir,
    column=PHOTO_COLUMN,
    first_row=2
):
    summary = {
        "total": len(rows),
        "embedded": 0,
        "local": 0,
        "external": 0,
        "empty": 0,
        "missing": [],
        "failed": []
    }

    if not PILLOW_AVAILABLE:
        logging.warning(
            "[FOTO] Pillow belum terinstall."
        )
        return summary

    image_dir = Path(image_dir)

    # Index foto lokal
    index = build_photo_index(image_dir)

    # Hapus gambar lama supaya tidak double
    worksheet._images = []

    # Atur lebar kolom
    worksheet.column_dimensions[
        get_column_letter(column)
    ].width = PHOTO_COLUMN_WIDTH

    for i, row in enumerate(rows, start=first_row):

        value = (row.get("visit_image") or "").strip()

        cell = worksheet.cell(
            row=i,
            column=column
        )

        # Tidak ada foto
        if not value:
            summary["empty"] += 1
            continue

        # ==========================================
        # 1. COBA CARI FOTO LOKAL
        # ==========================================

        photo_path = find_photo(
            value,
            index,
            row.get("id")
        )

        if photo_path is not None:
            try:
                image_stream = _prepare_photo(photo_path)

                excel_image = XLImage(image_stream)

                display_width, display_height = _display_size(
                    excel_image.width,
                    excel_image.height
                )

                excel_image.width = display_width
                excel_image.height = display_height

                worksheet.add_image(
                    excel_image,
                    cell.coordinate
                )

                cell.value = None

                worksheet.row_dimensions[i].height = PHOTO_ROW_HEIGHT

                summary["embedded"] += 1
                summary["local"] += 1

                continue

            except Exception as exc:
                logging.warning(
                    "[FOTO] Gagal memproses foto lokal %s: %s",
                    photo_path,
                    exc
                )

        # ==========================================
        # 2. KALAU LOKAL TIDAK ADA,
        #    COBA DOWNLOAD DARI GALENIUM
        # ==========================================

        external_url = build_external_url(value)

        if external_url:

            try:
                logging.info(
                    "[FOTO] Download dari URL: %s",
                    external_url
                )

                image_bytes = download_external_photo(
                    external_url
                )

                image_stream = _prepare_photo_bytes(
                    image_bytes
                )

                excel_image = XLImage(image_stream)

                display_width, display_height = _display_size(
                    excel_image.width,
                    excel_image.height
                )

                excel_image.width = display_width
                excel_image.height = display_height

                worksheet.add_image(
                    excel_image,
                    cell.coordinate
                )

                # Foto sudah berhasil ditempel,
                # URL tidak perlu ditampilkan di cell.
                cell.value = None

                worksheet.row_dimensions[i].height = PHOTO_ROW_HEIGHT

                summary["embedded"] += 1
                summary["external"] += 1

                continue

            except Exception as exc:

                logging.warning(
                    "[FOTO] Gagal download %s: %s",
                    external_url,
                    exc
                )

                # Kalau gagal download,
                # tetap simpan URL di Excel
                cell.value = external_url

                summary["failed"].append({
                    "value": value,
                    "url": external_url,
                    "error": str(exc)
                })

                continue

        # ==========================================
        # 3. BENAR-BENAR TIDAK ADA
        # ==========================================

        summary["missing"].append(value)

    logging.info(
        "[FOTO] %s dari %s baris berhasil ditempel foto.",
        summary["embedded"],
        summary["total"]
    )

    logging.info(
        "[FOTO] Lokal: %s | External: %s",
        summary["local"],
        summary["external"]
    )

    if summary["missing"]:
        logging.warning(
            "[FOTO] %s file tidak ditemukan.",
            len(summary["missing"])
        )

    if summary["failed"]:
        logging.warning(
            "[FOTO] %s foto gagal didownload.",
            len(summary["failed"])
        )

    return summary