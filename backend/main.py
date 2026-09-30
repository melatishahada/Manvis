from fastapi import FastAPI, Query, HTTPException
from pydantic import BaseModel
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse, JSONResponse, FileResponse

from sqlalchemy import text
from passlib.context import CryptContext
from database import engine

from io import BytesIO, StringIO
from pathlib import Path
from datetime import datetime
import csv

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="MANVIS API",
    description="Manual Visit Information System",
    version="1.0.0",
)

pwd_context = CryptContext(
    schemes=["bcrypt"],
    deprecated="auto"
)

class LoginRequest(BaseModel):
    username: str
    password: str
# =========================================================
# CORS
# =========================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.post("/login")
def login(data: LoginRequest):

    query = text("""
        SELECT
            id,
            name,
            email,
            password,
            role,
            is_active
        FROM login_users
        WHERE email = :username
        LIMIT 1
    """)

    with engine.connect() as conn:
        result = conn.execute(
            query,
            {
                "username": data.username
            }
        ).mappings().first()

    if not result:
        raise HTTPException(
            status_code=401,
            detail="Username atau password salah"
        )

    if result["is_active"] != 1:
        raise HTTPException(
            status_code=403,
            detail="Akun tidak aktif"
        )

    if not pwd_context.verify(
        data.password,
        result["password"]
    ):
        raise HTTPException(
            status_code=401,
            detail="Username atau password salah"
        )

    return {
        "message": "Login berhasil",
        "user": {
            "id": result["id"],
            "name": result["name"],
            "email": result["email"],
            "role": result["role"]
        }
    }


class RegisterRequest(BaseModel):
    name: str
    email: str
    password: str
    role: str = "USER"


@app.post("/register")
def register(data: RegisterRequest):

    if data.role not in ["ADMIN", "USER"]:
        raise HTTPException(
            status_code=400,
            detail="Role harus ADMIN atau USER."
        )

    check_query = text("""
        SELECT id
        FROM login_users
        WHERE email = :email
        LIMIT 1
    """)

    with engine.connect() as conn:
        existing = conn.execute(
            check_query,
            {"email": data.email}
        ).first()

    if existing:
        raise HTTPException(
            status_code=400,
            detail="Email sudah terdaftar."
        )

    hashed_password = pwd_context.hash(data.password)

    insert_query = text("""
        INSERT INTO login_users
        (name, email, password, role, is_active)
        VALUES
        (:name, :email, :password, :role, 1)
    """)

    with engine.begin() as conn:
        result = conn.execute(
            insert_query,
            {
                "name": data.name,
                "email": data.email,
                "password": hashed_password,
                "role": data.role,
            }
        )
        new_id = result.lastrowid

    return {
        "status": "success",
        "id": new_id,
        "name": data.name,
        "email": data.email,
        "role": data.role,
    }

@app.get("/users")
def get_users():

    query = text("""
        SELECT id, name, email, role, is_active, created_at
        FROM login_users
        ORDER BY id ASC
    """)

    with engine.connect() as conn:
        rows = conn.execute(query).mappings().all()

    data = []

    for row in rows:
        data.append({
            "id": row["id"],
            "name": row["name"],
            "email": row["email"],
            "role": row["role"],
            "isActive": row["is_active"],
            "createdAt": (
                row["created_at"].isoformat()
                if row["created_at"]
                else None
            ),
        })

    return {"total": len(data), "data": data}


class UpdateUserRequest(BaseModel):
    role: str | None = None
    is_active: int | None = None


@app.put("/users/{user_id}")
def update_user(user_id: int, data: UpdateUserRequest):

    if data.role is not None and data.role not in ["ADMIN", "USER"]:
        raise HTTPException(
            status_code=400,
            detail="Role harus ADMIN atau USER."
        )

    check_query = text("SELECT id FROM login_users WHERE id = :id LIMIT 1")

    with engine.connect() as conn:
        existing = conn.execute(check_query, {"id": user_id}).first()

    if not existing:
        raise HTTPException(status_code=404, detail="User tidak ditemukan.")

    updates = []
    params = {"id": user_id}

    if data.role is not None:
        updates.append("role = :role")
        params["role"] = data.role

    if data.is_active is not None:
        updates.append("is_active = :is_active")
        params["is_active"] = data.is_active

    if not updates:
        raise HTTPException(status_code=400, detail="Tidak ada data yang diubah.")

    update_query = text(
        f"UPDATE login_users SET {', '.join(updates)} WHERE id = :id"
    )

    with engine.begin() as conn:
        conn.execute(update_query, params)

    return {"status": "success", "id": user_id}


@app.delete("/users/{user_id}")
def delete_user(user_id: int):

    check_query = text("SELECT id FROM login_users WHERE id = :id LIMIT 1")

    with engine.connect() as conn:
        existing = conn.execute(check_query, {"id": user_id}).first()

    if not existing:
        raise HTTPException(status_code=404, detail="User tidak ditemukan.")

    delete_query = text("DELETE FROM login_users WHERE id = :id")

    with engine.begin() as conn:
        conn.execute(delete_query, {"id": user_id})

    return {"status": "success", "id": user_id}
# =========================================================
# PATH
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
IMAGE_DIR = BASE_DIR / "images"

FRONTEND_DIST = BASE_DIR.parent / "frontend" / "dist"
# =========================================================
# HELPER
# =========================================================

def get_manual_visit_data(start_date=None, end_date=None):
    query = text(
        """
        SELECT
    manual_visit.id,
    sales.name AS user_name,
    kowil.name AS kowil_code,
    outlet.name AS outlet_name,
    manual_visit.visit_date,
    manual_visit.visit_catatan,
    manual_visit.visit_status,
    manual_visit.visit_image,
    manual_visit.created_at,
    manual_visit.updated_at
FROM manual_visit
JOIN sales
    ON manual_visit.user_id = sales.id
JOIN kowil
    ON manual_visit.kowil_id = kowil.id
JOIN outlet
    ON manual_visit.outlet_id = outlet.id
        WHERE
            (:start_date IS NULL
                OR manual_visit.visit_date >= :start_date)
            AND
            (:end_date IS NULL
                OR manual_visit.visit_date <= :end_date)
        ORDER BY
            manual_visit.visit_date ASC,
            manual_visit.id ASC
        """
    )

    with engine.connect() as connection:

        result = connection.execute(
            query,
            {
                "start_date": start_date,
                "end_date": end_date,
            },
        )

        rows = result.mappings().all()

    return rows


def validate_dates(start_date, end_date):

    if start_date and end_date:

        if start_date > end_date:
            raise HTTPException(
                status_code=400,
                detail="start_date tidak boleh lebih besar dari end_date.",
            )

# =========================================================
# TEST DATABASE
# =========================================================

@app.get("/test-db")
def test_db():

    try:

        with engine.connect() as connection:

            result = connection.execute(
                text("SELECT DATABASE() AS database_name")
            )

            row = result.mappings().first()

        return {
            "status": "success",
            "database": row["database_name"],
        }

    except Exception as error:

        return {
            "status": "error",
            "message": str(error),
        }


# =========================================================
# TABLES
# =========================================================

@app.get("/tables")
def get_tables():

    query = text(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = DATABASE()
        ORDER BY table_name
        """
    )

    with engine.connect() as connection:

        result = connection.execute(query)

        tables = [
            row[0]
            for row in result.fetchall()
        ]

    return {
        "total": len(tables),
        "tables": tables,
    }


# =========================================================
# MANUAL VISITS
# =========================================================

@app.get("/manual-visits")
def manual_visits(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    data = []

    for row in rows:

        data.append(
            {
                "id": row["id"],
                "user_name": row["user_name"],
                "kowil_name": row["kowil_code"],
                "outlet_name": row["outlet_name"],
                "visit_date": (
                    row["visit_date"].isoformat()
                    if row["visit_date"]
                    else None
                ),
                "visit_catatan": row["visit_catatan"],
                "visit_status": row["visit_status"],
                "visit_image": row["visit_image"],
                "created_at": (
                    row["created_at"].isoformat(sep=" ")
                    if row["created_at"]
                    else None
                ),
                "updated_at": (
                    row["updated_at"].isoformat(sep=" ")
                    if row["updated_at"]
                    else None
                ),
            }
        )

    return {
        "total": len(data),
        "data": data,
    }


# =========================================================
# EXCEL
# =========================================================

@app.get("/reports/manual-visits/excel")
def download_excel(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Manual Visit"

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    headers = [
        "ID",
        "NAME",
        "KOWIL_CODE",
        "OUTLET_NAME",
        "VISIT_DATE",
        "VISIT_CATATAN",
        "VISIT_STATUS",
        "GAMBAR",
        "CREATED_AT",
        "UPDATED_AT",
    ]

    worksheet.append(headers)

    # Header style
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1D4ED8",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for cell in worksheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin_border

    worksheet.row_dimensions[1].height = 25

    # -----------------------------------------------------
    # DATA
    # -----------------------------------------------------

    for index, row in enumerate(rows, start=2):

        visit_date = row["visit_date"]

        if visit_date:
            visit_date_value = visit_date.strftime(
                "%Y-%m-%d"
            )
        else:
            visit_date_value = ""

        created_at_value = row["created_at"]
        updated_at_value = row["updated_at"]

        worksheet.cell(
            row=index,
            column=1,
            value=row["id"],
        )

        worksheet.cell(
            row=index,
            column=2,
            value=row["user_name"],
        )

        worksheet.cell(
            row=index,
            column=3,
            value=row["kowil_code"],
        )

        worksheet.cell(
            row=index,
            column=4,
            value=row["outlet_name"],
        )

        worksheet.cell(
            row=index,
            column=5,
            value=visit_date_value,
        )

        worksheet.cell(
            row=index,
            column=6,
            value=row["visit_catatan"],
        )

        worksheet.cell(
            row=index,
            column=7,
            value=row["visit_status"],
        )

        # Kolom GAMBAR
        image_cell = worksheet.cell(
            row=index,
            column=8,
        )

        visit_image = row["visit_image"] or ""

        # Tetap tulis path/nama file di kolom GAMBAR.
        # Jika file tersedia, gambar juga ditempelkan ke Excel.
        image_cell.value = visit_image

        if visit_image:
            image_path = BASE_DIR / visit_image

            if image_path.exists() and image_path.is_file():
                try:
                    excel_image = XLImage(str(image_path))

                    # Ukuran gambar di Excel
                    excel_image.width = 140
                    excel_image.height = 100

                    worksheet.add_image(
                        excel_image,
                        f"H{index}",
                    )

                    worksheet.row_dimensions[index].height = 80

                except Exception as image_error:
                    image_cell.value = (
                        f"{visit_image} | "
                        f"Gagal embed: {image_error}"
                    )

        created_cell = worksheet.cell(
            row=index,
            column=9,
            value=created_at_value,
        )
        if created_at_value:
            created_cell.number_format = "yyyy-mm-dd hh:mm:ss"

        updated_cell = worksheet.cell(
            row=index,
            column=10,
            value=updated_at_value,
        )
        if updated_at_value:
            updated_cell.number_format = "yyyy-mm-dd hh:mm:ss"

        # Border
        for column in range(1, 11):

            cell = worksheet.cell(
                row=index,
                column=column,
            )

            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center",
            )

    # -----------------------------------------------------
    # COLUMN WIDTH
    # -----------------------------------------------------

    widths = {
        "A": 10,
        "B": 22,
        "C": 20,
        "D": 35,
        "E": 15,
        "F": 45,
        "G": 15,
        "H": 22,
        "I": 24,
        "J": 24,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.freeze_panes = "A2"

    # -----------------------------------------------------
    # SAVE TO MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
    f"MANUAL_VISIT {start_date} - {end_date}.xlsx"
)

    headers_response = {
        "Content-Disposition": (
            f'attachment; filename="{filename}"'
        )
    }

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers=headers_response,
    )


# =========================================================
# CSV
# =========================================================

@app.get("/reports/manual-visits/csv")
def download_csv(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    output = StringIO()

    writer = csv.writer(output)

    writer.writerow(
        [
            "ID",
            "NAME",
            "KOWIL_CODE",
            "OUTLET_NAME",
            "VISIT_DATE",
            "VISIT_CATATAN",
            "VISIT_STATUS",
            "GAMBAR",
            "CREATED_AT",
            "UPDATED_AT",
        ]
    )

    for row in rows:

        writer.writerow(
            [
                row["id"],
                row["user_name"],
                row["kowil_code"],
                row["outlet_name"],
                (
                    row["visit_date"].strftime(
                        "%Y-%m-%d"
                    )
                    if row["visit_date"]
                    else ""
                ),
                row["visit_catatan"],
                row["visit_status"],
                row["visit_image"] or "",
                (
                    str(row["created_at"])
                    if row["created_at"]
                    else ""
                ),
                (
                    str(row["updated_at"])
                    if row["updated_at"]
                    else ""
                ),
            ]
        )

    csv_content = output.getvalue()

    filename = (
        f"MANUAL_VISIT {start_date} - {end_date}.csv"
    )

    return StreamingResponse(
        iter([csv_content.encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# =========================================================
# REPORT HISTORY
# =========================================================

@app.get("/reports/history")
def get_report_history():

    history_query = text(
        """
        SELECT
            id,
            filename,
            start_date,
            end_date,
            report_type,
            created_at
        FROM report_history
        ORDER BY created_at DESC, id DESC
        """
    )

    downloads_query = text(
        """
        SELECT
            report_downloads.report_id,
            login_users.id AS user_id,
            login_users.name AS user_name,
            report_downloads.downloaded_at
        FROM report_downloads
        JOIN login_users
            ON report_downloads.user_id = login_users.id
        ORDER BY
            report_downloads.downloaded_at ASC,
            report_downloads.id ASC
        """
    )

    with engine.connect() as connection:

        history_rows = connection.execute(
            history_query
        ).mappings().all()

        download_rows = connection.execute(
            downloads_query
        ).mappings().all()

    downloads_by_report = {}

    for row in download_rows:

        report_id = row["report_id"]

        if report_id not in downloads_by_report:
            downloads_by_report[report_id] = []

        downloads_by_report[report_id].append(
            {
                "userId": row["user_id"],
                "userName": row["user_name"],
                "downloadedAt": (
                    row["downloaded_at"].isoformat()
                    if row["downloaded_at"]
                    else None
                ),
            }
        )

    data = []

    for row in history_rows:

        data.append(
            {
                "id": row["id"],
                "filename": row["filename"],
                "startDate": (
                    row["start_date"].isoformat()
                    if row["start_date"]
                    else None
                ),
                "endDate": (
                    row["end_date"].isoformat()
                    if row["end_date"]
                    else None
                ),
                "type": row["report_type"],
                "createdAt": (
                    row["created_at"].isoformat()
                    if row["created_at"]
                    else None
                ),
                "downloadedBy": downloads_by_report.get(
                    row["id"],
                    []
                ),
            }
        )

    return {
        "total": len(data),
        "data": data,
    }

@app.get("/reports/history/{report_id}/download")
def download_report(
    report_id: int,
    user_id: int = Query(...)
):

    report_query = text(
        """
        SELECT
            id,
            filename,
            start_date,
            end_date,
            report_type
        FROM report_history
        WHERE id = :report_id
        LIMIT 1
        """
    )

    user_query = text(
        """
        SELECT
            id,
            name,
            email,
            role,
            is_active
        FROM login_users
        WHERE id = :user_id
        LIMIT 1
        """
    )

    with engine.connect() as connection:

        report = connection.execute(
            report_query,
            {
                "report_id": report_id
            }
        ).mappings().first()

        user = connection.execute(
            user_query,
            {
                "user_id": user_id
            }
        ).mappings().first()

    if not report:
        raise HTTPException(
            status_code=404,
            detail="Report tidak ditemukan."
        )

    if not user:
        raise HTTPException(
            status_code=404,
            detail="User tidak ditemukan."
        )

    if user["is_active"] != 1:
        raise HTTPException(
            status_code=403,
            detail="User tidak aktif."
        )

    start_date = report["start_date"]
    end_date = report["end_date"]
    report_type = report["report_type"]

    # Catat siapa yang melakukan download
    insert_download = text(
        """
        INSERT INTO report_downloads
        (
            report_id,
            user_id
        )
        VALUES
        (
            :report_id,
            :user_id
        )
        """
    )

    with engine.begin() as connection:
        connection.execute(
            insert_download,
            {
                "report_id": report_id,
                "user_id": user_id,
            }
        )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    # =========================
    # EXCEL
    # =========================

    if report_type == "Excel":

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Manual Visit"

        headers = [
            "ID",
            "NAME",
            "KOWIL_CODE",
            "OUTLET_NAME",
            "VISIT_DATE",
            "VISIT_CATATAN",
            "VISIT_STATUS",
            "GAMBAR",
            "CREATED_AT",
            "UPDATED_AT",
        ]

        worksheet.append(headers)

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1D4ED8",
        )

        header_font = Font(
            color="FFFFFF",
            bold=True,
        )

        thin_border = Border(
            left=Side(
                style="thin",
                color="D1D5DB"
            ),
            right=Side(
                style="thin",
                color="D1D5DB"
            ),
            top=Side(
                style="thin",
                color="D1D5DB"
            ),
            bottom=Side(
                style="thin",
                color="D1D5DB"
            ),
        )

        for cell in worksheet[1]:

            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )
            cell.border = thin_border

        worksheet.row_dimensions[1].height = 25

        for index, row in enumerate(
            rows,
            start=2
        ):

            visit_date = row["visit_date"]

            visit_date_value = (
                visit_date.strftime("%Y-%m-%d")
                if visit_date
                else ""
            )

            worksheet.cell(
                row=index,
                column=1,
                value=row["id"],
            )

            worksheet.cell(
                row=index,
                column=2,
                value=row["user_name"],
            )

            worksheet.cell(
                row=index,
                column=3,
                value=row["kowil_code"],
            )

            worksheet.cell(
                row=index,
                column=4,
                value=row["outlet_name"],
            )

            worksheet.cell(
                row=index,
                column=5,
                value=visit_date_value,
            )

            worksheet.cell(
                row=index,
                column=6,
                value=row["visit_catatan"],
            )

            worksheet.cell(
                row=index,
                column=7,
                value=row["visit_status"],
            )

            image_cell = worksheet.cell(
                row=index,
                column=8,
            )

            visit_image = row["visit_image"] or ""

            image_cell.value = visit_image

            if visit_image:

                image_path = BASE_DIR / visit_image

                if (
                    image_path.exists()
                    and image_path.is_file()
                ):

                    try:

                        excel_image = XLImage(
                            str(image_path)
                        )

                        excel_image.width = 140
                        excel_image.height = 100

                        worksheet.add_image(
                            excel_image,
                            f"H{index}",
                        )

                        worksheet.row_dimensions[
                            index
                        ].height = 80

                    except Exception as image_error:

                        image_cell.value = (
                            f"{visit_image} | "
                            f"Gagal embed: "
                            f"{image_error}"
                        )

            created_cell = worksheet.cell(
                row=index,
                column=9,
                value=row["created_at"],
            )

            if row["created_at"]:
                created_cell.number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )

            updated_cell = worksheet.cell(
                row=index,
                column=10,
                value=row["updated_at"],
            )

            if row["updated_at"]:
                updated_cell.number_format = (
                    "yyyy-mm-dd hh:mm:ss"
                )

            for column in range(1, 11):

                cell = worksheet.cell(
                    row=index,
                    column=column,
                )

                cell.border = thin_border

                cell.alignment = Alignment(
                    vertical="center",
                )

        widths = {
            "A": 10,
            "B": 22,
            "C": 20,
            "D": 35,
            "E": 15,
            "F": 45,
            "G": 15,
            "H": 22,
            "I": 24,
            "J": 24,
        }

        for column, width in widths.items():

            worksheet.column_dimensions[
                column
            ].width = width

        worksheet.freeze_panes = "A2"

        output = BytesIO()

        workbook.save(output)

        output.seek(0)

        return StreamingResponse(
            output,
            media_type=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            headers={
                "Content-Disposition": (
                    f'attachment; filename="{report["filename"]}"'
                )
            },
        )

    # =========================
    # CSV
    # =========================

    output = StringIO()

    writer = csv.writer(output)

    writer.writerow(
        [
            "ID",
            "NAME",
            "KOWIL_CODE",
            "OUTLET_NAME",
            "VISIT_DATE",
            "VISIT_CATATAN",
            "VISIT_STATUS",
            "GAMBAR",
            "CREATED_AT",
            "UPDATED_AT",
        ]
    )

    for row in rows:

        writer.writerow(
            [
                row["id"],
                row["user_name"],
                row["kowil_code"],
                row["outlet_name"],
                (
                    row["visit_date"].strftime(
                        "%Y-%m-%d"
                    )
                    if row["visit_date"]
                    else ""
                ),
                row["visit_catatan"],
                row["visit_status"],
                row["visit_image"] or "",
                (
                    str(row["created_at"])
                    if row["created_at"]
                    else ""
                ),
                (
                    str(row["updated_at"])
                    if row["updated_at"]
                    else ""
                ),
            ]
        )

    csv_content = output.getvalue()

    return StreamingResponse(
        iter(
            [
                csv_content.encode(
                    "utf-8-sig"
                )
            ]
        ),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{report["filename"]}"'
            )
        },
    )

    

# =========================================================
# CREATE REPORT HISTORY
# =========================================================

@app.post("/reports/history")
def create_report_history(payload: dict):

    start_date = payload.get("start_date")
    end_date = payload.get("end_date")
    report_type = payload.get("report_type")

    if not start_date or not end_date:

        raise HTTPException(
            status_code=400,
            detail="Tanggal awal dan akhir wajib diisi.",
        )

    if report_type not in ["Excel", "CSV"]:

        raise HTTPException(
            status_code=400,
            detail="report_type harus Excel atau CSV.",
        )

    validate_dates(
        start_date,
        end_date,
    )

    extension = (
        "xlsx"
        if report_type == "Excel"
        else "csv"
    )

    filename = (
        f"MANVIS {start_date} - "
        f"{end_date}.{extension}"
    )

    query = text(
        """
        INSERT INTO report_history
        (
            filename,
            start_date,
            end_date,
            report_type
        )
        VALUES
        (
            :filename,
            :start_date,
            :end_date,
            :report_type
        )
        """
    )

    with engine.begin() as connection:

        result = connection.execute(
            query,
            {
                "filename": filename,
                "start_date": start_date,
                "end_date": end_date,
                "report_type": report_type,
            },
        )

        report_id = result.lastrowid

    return {
        "status": "success",
        "id": report_id,
        "filename": filename,
        "startDate": start_date,
        "endDate": end_date,
        "type": report_type,
    }

# =========================================================
# SERVE REACT BUILD
# =========================================================

if not FRONTEND_DIST.exists():

    print(
        f"WARNING: React dist tidak ditemukan: "
        f"{FRONTEND_DIST}"
    )

else:

    app.mount(
        "/assets",
        StaticFiles(
            directory=FRONTEND_DIST / "assets"
        ),
        name="assets",
    )


    @app.get("/{full_path:path}")
    async def serve_react(full_path: str):

        requested_file = FRONTEND_DIST / full_path

        if (
            full_path
            and requested_file.exists()
            and requested_file.is_file()
        ):
            return FileResponse(
                requested_file
            )

        return FileResponse(
            FRONTEND_DIST / "index.html"
        )