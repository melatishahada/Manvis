from fastapi import FastAPI, Query, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import StreamingResponse, JSONResponse, FileResponse

from sqlalchemy import text
from database import engine

from io import BytesIO, StringIO
from pathlib import Path
from datetime import datetime
import csv

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# =========================================================
# APP
# =========================================================

app = FastAPI(
    title="MANVIS API",
    description="Manual Visit Information System",
    version="1.0.0",
)


# =========================================================
# CORS
# =========================================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://localhost:5174",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# =========================================================
# PATH
# =========================================================

BASE_DIR = Path(__file__).resolve().parent
IMAGE_DIR = BASE_DIR / "images"

FRONTEND_DIST = BASE_DIR.parent / "frontend" / "dist"
# =========================================================
# HELPER
# =========================================================

def get_manual_visit_data(start_date=None, end_date=None):
    query = text(
        """
        SELECT
            manual_visit.id,
            users.name AS user_name,
            kowil.name AS kowil_code,
            outlet.name AS outlet_name,
            manual_visit.visit_date,
            manual_visit.visit_catatan,
            manual_visit.visit_status,
            manual_visit.visit_image,
            manual_visit.created_at,
            manual_visit.updated_at
        FROM manual_visit
        JOIN users
            ON manual_visit.user_id = users.id
        JOIN kowil
            ON manual_visit.kowil_id = kowil.id
        JOIN outlet
            ON manual_visit.outlet_id = outlet.id
        WHERE
            (:start_date IS NULL
                OR manual_visit.visit_date >= :start_date)
            AND
            (:end_date IS NULL
                OR manual_visit.visit_date <= :end_date)
        ORDER BY
            manual_visit.visit_date ASC,
            manual_visit.id ASC
        """
    )

    with engine.connect() as connection:

        result = connection.execute(
            query,
            {
                "start_date": start_date,
                "end_date": end_date,
            },
        )

        rows = result.mappings().all()

    return rows


def validate_dates(start_date, end_date):

    if start_date and end_date:

        if start_date > end_date:
            raise HTTPException(
                status_code=400,
                detail="start_date tidak boleh lebih besar dari end_date.",
            )

# =========================================================
# TEST DATABASE
# =========================================================

@app.get("/test-db")
def test_db():

    try:

        with engine.connect() as connection:

            result = connection.execute(
                text("SELECT DATABASE() AS database_name")
            )

            row = result.mappings().first()

        return {
            "status": "success",
            "database": row["database_name"],
        }

    except Exception as error:

        return {
            "status": "error",
            "message": str(error),
        }


# =========================================================
# TABLES
# =========================================================

@app.get("/tables")
def get_tables():

    query = text(
        """
        SELECT table_name
        FROM information_schema.tables
        WHERE table_schema = DATABASE()
        ORDER BY table_name
        """
    )

    with engine.connect() as connection:

        result = connection.execute(query)

        tables = [
            row[0]
            for row in result.fetchall()
        ]

    return {
        "total": len(tables),
        "tables": tables,
    }


# =========================================================
# MANUAL VISITS
# =========================================================

@app.get("/manual-visits")
def manual_visits(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    data = []

    for row in rows:

        data.append(
            {
                "id": row["id"],
                "user_name": row["user_name"],
                "kowil_name": row["kowil_code"],
                "outlet_name": row["outlet_name"],
                "visit_date": (
                    row["visit_date"].isoformat()
                    if row["visit_date"]
                    else None
                ),
                "visit_catatan": row["visit_catatan"],
                "visit_status": row["visit_status"],
                "visit_image": row["visit_image"],
            }
        )

    return {
        "total": len(data),
        "data": data,
    }


# =========================================================
# EXCEL
# =========================================================

@app.get("/reports/manual-visits/excel")
def download_excel(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Manual Visit"

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    headers = [
        "ID",
        "NAME",
        "KOWIL_CODE",
        "OUTLET_NAME",
        "VISIT_DATE",
        "VISIT_CATATAN",
        "VISIT_STATUS",
        "GAMBAR",
        "CREATED_AT",
        "UPDATED_AT",
    ]

    worksheet.append(headers)

    # Header style
    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1D4ED8",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )

    for cell in worksheet[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin_border

    worksheet.row_dimensions[1].height = 25

    # -----------------------------------------------------
    # DATA
    # -----------------------------------------------------

    for index, row in enumerate(rows, start=2):

        visit_date = row["visit_date"]

        if visit_date:
            visit_date_value = visit_date.strftime(
                "%Y-%m-%d"
            )
        else:
            visit_date_value = ""

        created_at = row["created_at"]

        if created_at:
            created_at_value = str(created_at)
        else:
            created_at_value = ""

        updated_at = row["updated_at"]

        if updated_at:
            updated_at_value = str(updated_at)
        else:
            updated_at_value = ""

        worksheet.cell(
            row=index,
            column=1,
            value=row["id"],
        )

        worksheet.cell(
            row=index,
            column=2,
            value=row["user_name"],
        )

        worksheet.cell(
            row=index,
            column=3,
            value=row["kowil_code"],
        )

        worksheet.cell(
            row=index,
            column=4,
            value=row["outlet_name"],
        )

        worksheet.cell(
            row=index,
            column=5,
            value=visit_date_value,
        )

        worksheet.cell(
            row=index,
            column=6,
            value=row["visit_catatan"],
        )

        worksheet.cell(
            row=index,
            column=7,
            value=row["visit_status"],
        )

        # Kolom GAMBAR
        image_cell = worksheet.cell(
            row=index,
            column=8,
        )

        visit_image = row["visit_image"]

        if visit_image:

            image_path = BASE_DIR / visit_image

            if image_path.exists():

                try:

                    excel_image = XLImage(
                        str(image_path)
                    )

                    # Ukuran gambar
                    excel_image.width = 140
                    excel_image.height = 100

                    worksheet.add_image(
                        excel_image,
                        f"H{index}",
                    )

                    worksheet.row_dimensions[
                        index
                    ].height = 80

                except Exception as image_error:

                    image_cell.value = (
                        f"Gagal memasukkan gambar: "
                        f"{image_error}"
                    )

            else:

                image_cell.value = (
                    f"Gambar tidak ditemukan: "
                    f"{visit_image}"
                )

        worksheet.cell(
            row=index,
            column=9,
            value=created_at_value,
        )

        worksheet.cell(
            row=index,
            column=10,
            value=updated_at_value,
        )

        # Border
        for column in range(1, 11):

            cell = worksheet.cell(
                row=index,
                column=column,
            )

            cell.border = thin_border
            cell.alignment = Alignment(
                vertical="center",
            )

    # -----------------------------------------------------
    # COLUMN WIDTH
    # -----------------------------------------------------

    widths = {
        "A": 10,
        "B": 22,
        "C": 20,
        "D": 35,
        "E": 15,
        "F": 45,
        "G": 15,
        "H": 22,
        "I": 24,
        "J": 24,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    worksheet.freeze_panes = "A2"

    # -----------------------------------------------------
    # SAVE TO MEMORY
    # -----------------------------------------------------

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
    f"MANUAL_VISIT {start_date} - {end_date}.xlsx"
)

    headers_response = {
        "Content-Disposition": (
            f'attachment; filename="{filename}"'
        )
    }

    return StreamingResponse(
        output,
        media_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        ),
        headers=headers_response,
    )


# =========================================================
# CSV
# =========================================================

@app.get("/reports/manual-visits/csv")
def download_csv(
    start_date: str | None = Query(default=None),
    end_date: str | None = Query(default=None),
):

    validate_dates(
        start_date,
        end_date,
    )

    rows = get_manual_visit_data(
        start_date,
        end_date,
    )

    output = StringIO()

    writer = csv.writer(output)

    writer.writerow(
        [
            "ID",
            "NAME",
            "KOWIL_CODE",
            "OUTLET_NAME",
            "VISIT_DATE",
            "VISIT_CATATAN",
            "VISIT_STATUS",
            "GAMBAR",
            "CREATED_AT",
            "UPDATED_AT",
        ]
    )

    for row in rows:

        writer.writerow(
            [
                row["id"],
                row["user_name"],
                row["kowil_code"],
                row["outlet_name"],
                (
                    row["visit_date"].strftime(
                        "%Y-%m-%d"
                    )
                    if row["visit_date"]
                    else ""
                ),
                row["visit_catatan"],
                row["visit_status"],
                row["visit_image"] or "",
                (
                    str(row["created_at"])
                    if row["created_at"]
                    else ""
                ),
                (
                    str(row["updated_at"])
                    if row["updated_at"]
                    else ""
                ),
            ]
        )

    csv_content = output.getvalue()

    filename = (
        f"MANUAL_VISIT {start_date} - {end_date}.csv"
    )

    return StreamingResponse(
        iter([csv_content.encode("utf-8-sig")]),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# =========================================================
# REPORT HISTORY
# =========================================================

@app.get("/reports/history")
def get_report_history():

    query = text(
        """
        SELECT
            id,
            filename,
            start_date,
            end_date,
            report_type,
            created_at
        FROM report_history
        ORDER BY created_at DESC, id DESC
        """
    )

    with engine.connect() as connection:

        result = connection.execute(query)

        rows = result.mappings().all()

    data = []

    for row in rows:

        data.append(
            {
                "id": row["id"],
                "filename": row["filename"],
                "startDate": (
                    row["start_date"].isoformat()
                    if row["start_date"]
                    else None
                ),
                "endDate": (
                    row["end_date"].isoformat()
                    if row["end_date"]
                    else None
                ),
                "type": row["report_type"],
                "createdAt": (
                    row["created_at"].isoformat()
                    if row["created_at"]
                    else None
                ),
            }
        )

    return {
        "total": len(data),
        "data": data,
    }


# =========================================================
# CREATE REPORT HISTORY
# =========================================================

@app.post("/reports/history")
def create_report_history(payload: dict):

    start_date = payload.get("start_date")
    end_date = payload.get("end_date")
    report_type = payload.get("report_type")

    if not start_date or not end_date:

        raise HTTPException(
            status_code=400,
            detail="Tanggal awal dan akhir wajib diisi.",
        )

    if report_type not in ["Excel", "CSV"]:

        raise HTTPException(
            status_code=400,
            detail="report_type harus Excel atau CSV.",
        )

    validate_dates(
        start_date,
        end_date,
    )

    extension = (
        "xlsx"
        if report_type == "Excel"
        else "csv"
    )

    filename = (
        f"MANVIS {start_date} - "
        f"{end_date}.{extension}"
    )

    query = text(
        """
        INSERT INTO report_history
        (
            filename,
            start_date,
            end_date,
            report_type
        )
        VALUES
        (
            :filename,
            :start_date,
            :end_date,
            :report_type
        )
        """
    )

    with engine.begin() as connection:

        result = connection.execute(
            query,
            {
                "filename": filename,
                "start_date": start_date,
                "end_date": end_date,
                "report_type": report_type,
            },
        )

        report_id = result.lastrowid

    return {
        "status": "success",
        "id": report_id,
        "filename": filename,
        "startDate": start_date,
        "endDate": end_date,
        "type": report_type,
    }

# =========================================================
# SERVE REACT BUILD
# =========================================================

if not FRONTEND_DIST.exists():

    print(
        f"WARNING: React dist tidak ditemukan: "
        f"{FRONTEND_DIST}"
    )

else:

    app.mount(
        "/assets",
        StaticFiles(
            directory=FRONTEND_DIST / "assets"
        ),
        name="assets",
    )


    @app.get("/{full_path:path}")
    async def serve_react(full_path: str):

        requested_file = FRONTEND_DIST / full_path

        if (
            full_path
            and requested_file.exists()
            and requested_file.is_file()
        ):
            return FileResponse(
                requested_file
            )

        return FileResponse(
            FRONTEND_DIST / "index.html"
        )